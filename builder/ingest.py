"""Read a case folder (or a zip of one) into text documents for the model roles."""
import hashlib
import json
import pathlib
import tempfile
import zipfile

TEXT_SUFFIXES = {".md", ".txt", ".json", ".csv", ".tsv", ".yaml", ".yml", ".toml", ".xml", ".html", ".htm", ".js", ".ts", ".py",
                 ".drawio", ".eml", ".ini", ".rst", ".log", ""}
SKIP_DIRS = {".git", "node_modules", "__pycache__", ".wrangler"}
FIRST_STEMS = ("brief", "readme")
MAX_CHARS = 60000       # a longer file is cut; for JSON an outline of its structure is added
HEAD_CHARS = 12000


def _pdf(path):
    try:
        import fitz
        with fitz.open(str(path)) as doc:
            return "\n\n".join(f"[page {i + 1}]\n{page.get_text('text')}" for i, page in enumerate(doc))
    except ImportError:
        import pypdf
        reader = pypdf.PdfReader(str(path))
        return "\n\n".join(f"[page {i + 1}]\n{page.extract_text() or ''}" for i, page in enumerate(reader.pages))


def _docx(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = docx.Document(str(path))
    out = []
    for child in doc.element.body.iterchildren():      # paragraphs and tables in document order
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            style = (p.style.name or "") if p.style is not None else ""
            text = p.text.strip()
            if text:
                out.append(("#" * int(style[-1]) + " " if style.startswith("Heading") and style[-1].isdigit() else "") + text)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                out.append(" | ".join(" ".join(cell.text.split()) for cell in row.cells))
    return "\n".join(out)


def _xlsx(path):
    import openpyxl
    book = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    out = []
    for sheet in book.worksheets:
        out.append(f"[sheet {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            if any(c not in (None, "") for c in row):
                out.append(" | ".join("" if c is None else str(c) for c in row))
    return "\n".join(out)


READERS = {".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx, ".xlsm": _xlsx}


def _outline(value, depth=0):
    """The structure of a large JSON value: keys, list lengths and the shape of the first list item."""
    if depth > 6:
        return "..."
    if isinstance(value, dict):
        keys = list(value.items())
        shown = ", ".join(f"{json.dumps(k, ensure_ascii=False)}: {_outline(v, depth + 1)}" for k, v in keys[:12])
        return "{" + shown + (f", ... {len(keys) - 12} more keys" if len(keys) > 12 else "") + "}"
    if isinstance(value, list):
        return f"[{len(value)} items, first: {_outline(value[0], depth + 1)}]" if value else "[]"
    return type(value).__name__ if not isinstance(value, str) else "string"


def _cut(text, suffix):
    if len(text) <= MAX_CHARS:
        return text, ""
    note = f"cut to the first {HEAD_CHARS} of {len(text)} characters"
    extra = ""
    if suffix == ".json":
        try:
            extra = "\n\n[structure of the whole file]\n" + _outline(json.loads(text))
        except Exception:
            pass
    return text[:HEAD_CHARS] + "\n[... the file continues ...]" + extra, note


def read_folder(folder):
    """Every file of the folder as a document: brief and readme first, then by path. A file that cannot be read is listed with the reason."""
    root = pathlib.Path(folder)
    if root.is_file() and root.suffix.lower() == ".zip":
        target = pathlib.Path(tempfile.mkdtemp(prefix="case-"))
        with zipfile.ZipFile(root) as z:
            z.extractall(target)
        inner = [p for p in target.iterdir() if p.name != "__MACOSX"]
        root = inner[0] if len(inner) == 1 and inner[0].is_dir() else target
    if not root.is_dir():
        raise ValueError(f"the case folder does not exist: {folder}")
    files = [p for p in root.rglob("*") if p.is_file() and not (set(p.relative_to(root).parts) & SKIP_DIRS)]

    def order(p):
        rel = p.relative_to(root)
        first = len(rel.parts) == 1 and p.stem.lower() in FIRST_STEMS
        return (0 if first else 1, FIRST_STEMS.index(p.stem.lower()) if first else 0, rel.as_posix().lower())

    docs = []
    for p in sorted(files, key=order):
        rel, suffix = p.relative_to(root).as_posix(), p.suffix.lower()
        raw = p.read_bytes()
        doc = {"path": rel, "sha256": hashlib.sha256(raw.replace(b"\r\n", b"\n")).hexdigest(), "bytes": len(raw), "text": "", "note": ""}
        try:
            if suffix in READERS:
                text = READERS[suffix](p)
            elif suffix in TEXT_SUFFIXES:
                text = raw.decode("utf-8-sig", errors="replace")
            else:
                text, doc["note"] = "", "file type not read"
            if text:
                doc["text"], cut = _cut(text.replace("\r\n", "\n"), suffix)
                doc["note"] = cut
            elif not doc["note"]:
                doc["note"] = "no text found in the file"
        except Exception as e:
            doc["note"] = f"could not be read: {type(e).__name__}: {str(e)[:120]}"
        docs.append(doc)
    return docs


def documents_block(docs):
    out = []
    for d in docs:
        note = f' note="{d["note"]}"' if d["note"] else ""
        out.append(f'<document path="{d["path"]}"{note}>\n{d["text"]}\n</document>')
    return "<case_folder>\n" + "\n".join(out) + "\n</case_folder>"
